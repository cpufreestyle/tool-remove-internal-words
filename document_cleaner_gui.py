#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文档内部字样清除工具 - 跨平台 GUI 版本
支持: macOS, Windows, Linux
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import os
import sys
import platform
import threading
import queue
from datetime import datetime
from pathlib import Path
import re
import shutil

# 平台检测
PLATFORM = platform.system()
IS_MAC = PLATFORM == 'Darwin'
IS_WINDOWS = PLATFORM == 'Windows'
IS_LINUX = PLATFORM == 'Linux'

# 尝试导入可选依赖
try:
    from PIL import Image, ImageTk
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import pytesseract
    HAS_TESSERACT = True
except ImportError:
    HAS_TESSERACT = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import PyPDF2
    HAS_PYPDF2 = True
except ImportError:
    HAS_PYPDF2 = False


class FileItem:
    """文件项类"""
    def __init__(self, filepath):
        self.filepath = filepath
        self.filename = os.path.basename(filepath)
        self.size = os.path.getsize(filepath)
        self.status = 'pending'  # pending, to-delete, to-keep, processed
        self.selected = False
        self.content = ''
        self.processed_content = ''
        
    def get_size_str(self):
        """获取文件大小字符串"""
        if self.size < 1024:
            return f"{self.size} B"
        elif self.size < 1024 * 1024:
            return f"{self.size / 1024:.1f} KB"
        else:
            return f"{self.size / 1024 / 1024:.1f} MB"
    
    def get_icon(self):
        """获取文件图标"""
        ext = os.path.splitext(self.filename)[1].lower()
        icons = {
            '.pdf': '📕',
            '.doc': '📘',
            '.docx': '📘',
            '.txt': '📄',
            '.jpg': '🖼️',
            '.jpeg': '🖼️',
            '.png': '🖼️',
            '.gif': '🖼️',
            '.bmp': '🖼️',
        }
        return icons.get(ext, '📎')


class DocumentCleanerApp:
    """文档清除工具主应用"""
    
    def __init__(self, root):
        self.root = root
        self.root.title("文档内部字样清除工具 v2.0")
        
        # 根据平台设置窗口大小
        if IS_MAC:
            self.root.geometry("1200x800")
        elif IS_WINDOWS:
            self.root.geometry("1200x800")
            # Windows 高 DPI 支持
            try:
                from ctypes import windll
                windll.shcore.SetProcessDpiAwareness(1)
            except:
                pass
        else:
            self.root.geometry("1200x800")
        
        self.root.configure(bg='#f0f0f0')
        
        # 数据
        self.files = []
        self.replace_count = 0
        self.log_queue = queue.Queue()
        
        # 默认替换规则
        self.default_rules = {
            '内部': True,
            '机密': True,
            '秘密': True,
            '绝密': False,
        }
        
        # 创建界面
        self.create_ui()
        
        # 启动日志更新
        self.update_log()
        
        # 检查依赖
        self.check_dependencies()
        
    def check_dependencies(self):
        """检查依赖库"""
        missing = []
        if not HAS_PIL:
            missing.append("Pillow (图片处理)")
        if not HAS_TESSERACT:
            missing.append("pytesseract (OCR识别)")
        if not HAS_DOCX:
            missing.append("python-docx (Word文档)")
        if not HAS_PYPDF2:
            missing.append("PyPDF2 (PDF文档)")
        
        if missing:
            self.log(f"⚠️ 缺少依赖库: {', '.join(missing)}", "warning")
            self.log("💡 运行: pip install Pillow pytesseract python-docx PyPDF2", "info")
        else:
            self.log("✅ 所有依赖库已安装", "success")
        
        # 显示平台信息
        self.log(f"🖥️ 运行平台: {PLATFORM} ({platform.machine()})", "info")
    
    def create_ui(self):
        """创建用户界面"""
        # 主框架
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # 样式配置
        style = ttk.Style()
        
        # 根据平台选择字体
        if IS_MAC:
            font_family = 'PingFang SC'
        elif IS_WINDOWS:
            font_family = 'Microsoft YaHei'
        else:
            font_family = 'Noto Sans CJK SC'
        
        style.configure('Title.TLabel', font=(font_family, 16, 'bold'))
        style.configure('Header.TLabel', font=(font_family, 12, 'bold'))
        style.configure('Status.TLabel', font=(font_family, 10))
        
        # 标题
        title_frame = ttk.Frame(main_frame)
        title_frame.pack(fill=tk.X, pady=(0, 10))
        ttk.Label(title_frame, text="📄 文档内部字样清除工具", style='Title.TLabel').pack(side=tk.LEFT)
        
        # 平台标识
        platform_text = f"[{PLATFORM}]"
        ttk.Label(title_frame, text=platform_text, style='Status.TLabel').pack(side=tk.RIGHT)
        
        # 创建笔记本（标签页）
        notebook = ttk.Notebook(main_frame)
        notebook.pack(fill=tk.BOTH, expand=True)
        
        # 标签页1: 文件处理
        file_tab = ttk.Frame(notebook, padding="10")
        notebook.add(file_tab, text="📁 文件处理")
        self.create_file_tab(file_tab)
        
        # 标签页2: 设置
        settings_tab = ttk.Frame(notebook, padding="10")
        notebook.add(settings_tab, text="⚙️ 设置")
        self.create_settings_tab(settings_tab)
        
        # 标签页3: 日志
        log_tab = ttk.Frame(notebook, padding="10")
        notebook.add(log_tab, text="📋 日志")
        self.create_log_tab(log_tab)
        
        # 底部状态栏
        self.create_status_bar(main_frame)
        
    def create_file_tab(self, parent):
        """创建文件处理标签页"""
        # 左右分栏
        paned = ttk.PanedWindow(parent, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True)
        
        # 左侧：文件列表
        left_frame = ttk.LabelFrame(paned, text="文件列表", padding="10")
        paned.add(left_frame, weight=1)
        
        # 文件列表工具栏
        toolbar = ttk.Frame(left_frame)
        toolbar.pack(fill=tk.X, pady=(0, 5))
        
        ttk.Button(toolbar, text="➕ 添加文件", command=self.add_files).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="📁 添加文件夹", command=self.add_folder).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="🗑️ 清空列表", command=self.clear_files).pack(side=tk.LEFT, padx=2)
        
        # 文件列表
        list_frame = ttk.Frame(left_frame)
        list_frame.pack(fill=tk.BOTH, expand=True)
        
        # 创建 Treeview
        columns = ('select', 'name', 'size', 'status')
        self.file_tree = ttk.Treeview(list_frame, columns=columns, show='headings', height=15)
        
        self.file_tree.heading('select', text='☑')
        self.file_tree.heading('name', text='文件名')
        self.file_tree.heading('size', text='大小')
        self.file_tree.heading('status', text='状态')
        
        self.file_tree.column('select', width=40, anchor='center')
        self.file_tree.column('name', width=300)
        self.file_tree.column('size', width=80, anchor='e')
        self.file_tree.column('status', width=100, anchor='center')
        
        # 滚动条
        scrollbar = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.file_tree.yview)
        self.file_tree.configure(yscrollcommand=scrollbar.set)
        
        self.file_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # 绑定事件
        self.file_tree.bind('<Button-1>', self.on_tree_click)
        self.file_tree.bind('<Double-1>', self.on_tree_double_click)
        
        # 批量操作按钮
        batch_frame = ttk.LabelFrame(left_frame, text="批量操作", padding="5")
        batch_frame.pack(fill=tk.X, pady=(10, 0))
        
        ttk.Button(batch_frame, text="☑️ 全选", command=self.select_all).pack(side=tk.LEFT, padx=2)
        ttk.Button(batch_frame, text="⬜ 取消全选", command=self.deselect_all).pack(side=tk.LEFT, padx=2)
        ttk.Button(batch_frame, text="🗑️ 标记删除", command=self.mark_delete).pack(side=tk.LEFT, padx=2)
        ttk.Button(batch_frame, text="✅ 标记保留", command=self.mark_keep).pack(side=tk.LEFT, padx=2)
        ttk.Button(batch_frame, text="❌ 移除选中", command=self.remove_selected).pack(side=tk.LEFT, padx=2)
        
        # 右侧：文本处理
        right_frame = ttk.LabelFrame(paned, text="文本处理", padding="10")
        paned.add(right_frame, weight=1)
        
        # 原始文本
        ttk.Label(right_frame, text="原始文本:", style='Header.TLabel').pack(anchor=tk.W)
        self.input_text = scrolledtext.ScrolledText(right_frame, height=10, wrap=tk.WORD)
        self.input_text.pack(fill=tk.BOTH, expand=True, pady=(5, 10))
        
        # 处理按钮
        btn_frame = ttk.Frame(right_frame)
        btn_frame.pack(fill=tk.X, pady=5)
        
        ttk.Button(btn_frame, text="🚀 处理文本", command=self.process_text).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="🚀 处理选中文件", command=self.process_files).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="📋 复制结果", command=self.copy_result).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="💾 保存结果", command=self.save_result).pack(side=tk.LEFT, padx=5)
        
        # 处理后文本
        ttk.Label(right_frame, text="处理后文本:", style='Header.TLabel').pack(anchor=tk.W)
        self.output_text = scrolledtext.ScrolledText(right_frame, height=10, wrap=tk.WORD)
        self.output_text.pack(fill=tk.BOTH, expand=True, pady=(5, 0))
        
    def create_settings_tab(self, parent):
        """创建设置标签页"""
        # 替换规则
        rules_frame = ttk.LabelFrame(parent, text="替换规则", padding="10")
        rules_frame.pack(fill=tk.X, pady=5)
        
        self.rule_vars = {}
        for i, (word, default) in enumerate(self.default_rules.items()):
            var = tk.BooleanVar(value=default)
            self.rule_vars[word] = var
            cb = ttk.Checkbutton(rules_frame, text=f"去除 \"{word}\" 字样", variable=var)
            cb.grid(row=i // 2, column=i % 2, sticky=tk.W, padx=10, pady=5)
        
        # 自定义规则
        custom_frame = ttk.LabelFrame(parent, text="自定义规则（每行一个）", padding="10")
        custom_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        self.custom_rules_text = scrolledtext.ScrolledText(custom_frame, height=10)
        self.custom_rules_text.pack(fill=tk.BOTH, expand=True)
        self.custom_rules_text.insert(tk.END, "# 输入自定义要去除的文字，每行一个\n# 例如：\n# 仅限内部使用\n# 严禁外传")
        
        # 选项
        options_frame = ttk.LabelFrame(parent, text="处理选项", padding="10")
        options_frame.pack(fill=tk.X, pady=5)
        
        self.scan_images_var = tk.BooleanVar(value=True)
        self.add_watermark_var = tk.BooleanVar(value=True)
        self.backup_var = tk.BooleanVar(value=True)
        
        ttk.Checkbutton(options_frame, text="扫描图片中的文字 (需要 Tesseract OCR)", 
                       variable=self.scan_images_var).grid(row=0, column=0, sticky=tk.W, pady=5)
        ttk.Checkbutton(options_frame, text="添加处理水印", 
                       variable=self.add_watermark_var).grid(row=1, column=0, sticky=tk.W, pady=5)
        ttk.Checkbutton(options_frame, text="备份原文件", 
                       variable=self.backup_var).grid(row=2, column=0, sticky=tk.W, pady=5)
        
        # OCR 设置
        ocr_frame = ttk.LabelFrame(parent, text="OCR 设置", padding="10")
        ocr_frame.pack(fill=tk.X, pady=5)
        
        ttk.Label(ocr_frame, text="Tesseract 路径:").grid(row=0, column=0, sticky=tk.W, pady=5)
        self.tesseract_path = tk.StringVar(value=self.get_default_tesseract_path())
        ttk.Entry(ocr_frame, textvariable=self.tesseract_path, width=50).grid(row=0, column=1, padx=5, pady=5)
        ttk.Button(ocr_frame, text="浏览...", command=self.browse_tesseract).grid(row=0, column=2, padx=5, pady=5)
        
        ttk.Label(ocr_frame, text="OCR 语言:").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.ocr_lang = tk.StringVar(value='chi_sim+eng')
        ttk.Combobox(ocr_frame, textvariable=self.ocr_lang, 
                    values=['chi_sim', 'eng', 'chi_sim+eng'], width=20).grid(row=1, column=1, sticky=tk.W, padx=5, pady=5)
        
    def get_default_tesseract_path(self):
        """获取默认 Tesseract 路径"""
        if IS_MAC:
            return '/usr/local/bin/tesseract'
        elif IS_WINDOWS:
            return r'C:\Program Files\Tesseract-OCR\tesseract.exe'
        else:
            return '/usr/bin/tesseract'
        
    def create_log_tab(self, parent):
        """创建日志标签页"""
        self.log_text = scrolledtext.ScrolledText(parent, height=30, wrap=tk.WORD)
        self.log_text.pack(fill=tk.BOTH, expand=True)
        
        # 配置标签颜色
        self.log_text.tag_configure('info', foreground='#45b7d1')
        self.log_text.tag_configure('success', foreground='#4ecdc4')
        self.log_text.tag_configure('warning', foreground='#f39c12')
        self.log_text.tag_configure('error', foreground='#ff6b6b')
        
        # 清空按钮
        ttk.Button(parent, text="清空日志", command=self.clear_log).pack(pady=5)
        
    def create_status_bar(self, parent):
        """创建状态栏"""
        status_frame = ttk.Frame(parent)
        status_frame.pack(fill=tk.X, pady=(5, 0))
        
        self.status_var = tk.StringVar(value="就绪")
        ttk.Label(status_frame, textvariable=self.status_var, style='Status.TLabel').pack(side=tk.LEFT)
        
        # 统计信息
        self.stats_var = tk.StringVar(value="文件: 0 | 待删除: 0 | 待保留: 0 | 已处理: 0 | 替换: 0")
        ttk.Label(status_frame, textvariable=self.stats_var, style='Status.TLabel').pack(side=tk.RIGHT)
        
    def add_files(self):
        """添加文件"""
        files = filedialog.askopenfilenames(
            title="选择文件",
            filetypes=[
                ("所有文件", "*.*"),
                ("Word 文档", "*.docx *.doc"),
                ("PDF 文件", "*.pdf"),
                ("文本文件", "*.txt"),
                ("图片文件", "*.jpg *.jpeg *.png *.gif *.bmp"),
            ]
        )
        
        for filepath in files:
            if not any(f.filepath == filepath for f in self.files):
                self.files.append(FileItem(filepath))
                self.log(f"添加文件: {os.path.basename(filepath)}", "info")
        
        self.refresh_file_list()
        
    def add_folder(self):
        """添加文件夹"""
        folder = filedialog.askdirectory(title="选择文件夹")
        if not folder:
            return
            
        # 支持的扩展名
        extensions = {'.docx', '.doc', '.pdf', '.txt', '.jpg', '.jpeg', '.png', '.gif', '.bmp'}
        
        count = 0
        for root, dirs, files in os.walk(folder):
            for filename in files:
                ext = os.path.splitext(filename)[1].lower()
                if ext in extensions:
                    filepath = os.path.join(root, filename)
                    if not any(f.filepath == filepath for f in self.files):
                        self.files.append(FileItem(filepath))
                        count += 1
        
        self.log(f"从文件夹添加 {count} 个文件", "info")
        self.refresh_file_list()
        
    def clear_files(self):
        """清空文件列表"""
        if self.files and messagebox.askyesno("确认", "确定要清空文件列表吗？"):
            self.files = []
            self.refresh_file_list()
            self.log("已清空文件列表", "info")
            
    def refresh_file_list(self):
        """刷新文件列表"""
        # 清空现有项
        for item in self.file_tree.get_children():
            self.file_tree.delete(item)
        
        # 添加文件
        status_map = {
            'pending': '⏳ 待处理',
            'to-delete': '🗑️ 待删除',
            'to-keep': '✅ 待保留',
            'processed': '✓ 已处理',
        }
        
        for file in self.files:
            select_mark = '☑' if file.selected else '⬜'
            status_text = status_map.get(file.status, file.status)
            
            item_id = self.file_tree.insert('', tk.END, values=(
                select_mark,
                f"{file.get_icon()} {file.filename}",
                file.get_size_str(),
                status_text
            ))
            file._tree_id = item_id
        
        self.update_stats()
        
    def on_tree_click(self, event):
        """处理树形列表点击"""
        region = self.file_tree.identify('region', event.x, event.y)
        if region == 'cell':
            column = self.file_tree.identify_column(event.x)
            if column == '#1':  # 选择列
                item = self.file_tree.identify_row(event.y)
                if item:
                    idx = self.file_tree.index(item)
                    if idx < len(self.files):
                        self.files[idx].selected = not self.files[idx].selected
                        self.refresh_file_list()
                        
    def on_tree_double_click(self, event):
        """处理双击事件 - 预览文件"""
        item = self.file_tree.identify_row(event.y)
        if item:
            idx = self.file_tree.index(item)
            if idx < len(self.files):
                file = self.files[idx]
                self.preview_file(file)
                
    def preview_file(self, file):
        """预览文件内容"""
        try:
            if file.content:
                self.input_text.delete('1.0', tk.END)
                self.input_text.insert(tk.END, file.content)
                self.log(f"预览文件: {file.filename}", "info")
            else:
                # 尝试读取文件
                content = self.read_file_content(file.filepath)
                if content:
                    file.content = content
                    self.input_text.delete('1.0', tk.END)
                    self.input_text.insert(tk.END, content)
                    self.log(f"预览文件: {file.filename}", "info")
                else:
                    messagebox.showinfo("提示", "无法预览此文件类型")
        except Exception as e:
            self.log(f"预览失败: {e}", "error")
            
    def read_file_content(self, filepath):
        """读取文件内容"""
        ext = os.path.splitext(filepath)[1].lower()
        
        try:
            if ext == '.txt':
                # 尝试多种编码
                for encoding in ['utf-8', 'gbk', 'gb2312', 'latin-1']:
                    try:
                        with open(filepath, 'r', encoding=encoding) as f:
                            return f.read()
                    except UnicodeDecodeError:
                        continue
                return ''
                    
            elif ext == '.docx' and HAS_DOCX:
                doc = docx.Document(filepath)
                return '\n'.join([para.text for para in doc.paragraphs])
                
            elif ext == '.pdf' and HAS_PYPDF2:
                with open(filepath, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    text = []
                    for page in reader.pages:
                        text.append(page.extract_text())
                    return '\n'.join(text)
                    
            elif ext in {'.jpg', '.jpeg', '.png', '.gif', '.bmp'} and HAS_TESSERACT and HAS_PIL:
                return self.ocr_image(filepath)
                
        except Exception as e:
            self.log(f"读取文件失败: {e}", "error")
            
        return ''
        
    def ocr_image(self, filepath):
        """OCR 识别图片"""
        if not HAS_TESSERACT or not HAS_PIL:
            return ''
            
        try:
            # 设置 Tesseract 路径
            tesseract_cmd = self.tesseract_path.get()
            if tesseract_cmd and os.path.exists(tesseract_cmd):
                pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
            
            image = Image.open(filepath)
            text = pytesseract.image_to_string(image, lang=self.ocr_lang.get())
            return text
        except Exception as e:
            self.log(f"OCR 识别失败: {e}", "error")
            return ''
            
    def select_all(self):
        """全选"""
        for file in self.files:
            file.selected = True
        self.refresh_file_list()
        
    def deselect_all(self):
        """取消全选"""
        for file in self.files:
            file.selected = False
        self.refresh_file_list()
        
    def mark_delete(self):
        """标记删除"""
        count = sum(1 for f in self.files if f.selected)
        for file in self.files:
            if file.selected:
                file.status = 'to-delete'
        self.refresh_file_list()
        self.log(f"标记 {count} 个文件为删除", "warning")
        
    def mark_keep(self):
        """标记保留"""
        count = sum(1 for f in self.files if f.selected)
        for file in self.files:
            if file.selected:
                file.status = 'to-keep'
        self.refresh_file_list()
        self.log(f"标记 {count} 个文件为保留", "success")
        
    def remove_selected(self):
        """移除选中"""
        if messagebox.askyesno("确认", "确定要移除选中的文件吗？"):
            self.files = [f for f in self.files if not f.selected]
            self.refresh_file_list()
            self.log("已移除选中文件", "info")
            
    def process_text(self):
        """处理文本"""
        text = self.input_text.get('1.0', tk.END).strip()
        if not text:
            messagebox.showwarning("警告", "请输入要处理的文本！")
            return
            
        processed, count = self.apply_rules(text)
        
        # 添加水印
        if self.add_watermark_var.get():
            timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            processed = f"【本文档已脱敏处理】\n处理时间：{timestamp}\n\n{processed}"
        
        self.output_text.delete('1.0', tk.END)
        self.output_text.insert(tk.END, processed)
        
        self.replace_count += count
        self.update_stats()
        self.log(f"处理完成，替换 {count} 处", "success")
        
    def process_files(self):
        """处理选中文件"""
        selected = [f for f in self.files if f.selected]
        if not selected:
            messagebox.showwarning("警告", "请先选择文件！")
            return
            
        self.log(f"开始处理 {len(selected)} 个文件...", "info")
        
        total_replaced = 0
        for file in selected:
            if file.status == 'to-delete':
                self.log(f"跳过标记删除的文件: {file.filename}", "warning")
                continue
                
            # 读取内容
            if not file.content:
                file.content = self.read_file_content(file.filepath)
                
            if file.content:
                # 处理
                processed, count = self.apply_rules(file.content)
                file.processed_content = processed
                file.status = 'processed'
                total_replaced += count
                self.log(f"处理完成: {file.filename}，替换 {count} 处", "success")
            else:
                self.log(f"无法处理: {file.filename}", "error")
                
        self.replace_count += total_replaced
        self.refresh_file_list()
        self.log(f"全部处理完成，共替换 {total_replaced} 处", "success")
        
    def apply_rules(self, text):
        """应用替换规则"""
        count = 0
        
        # 默认规则
        for word, var in self.rule_vars.items():
            if var.get():
                patterns = [
                    word,
                    f"[{word}]",
                    f"（{word}）",
                    f"【{word}】",
                ]
                for pattern in patterns:
                    matches = len(re.findall(re.escape(pattern), text))
                    if matches:
                        count += matches
                        text = text.replace(pattern, '')
        
        # 自定义规则
        custom_rules = self.custom_rules_text.get('1.0', tk.END).strip()
        for line in custom_rules.split('\n'):
            line = line.strip()
            if line and not line.startswith('#'):
                matches = len(re.findall(re.escape(line), text))
                if matches:
                    count += matches
                    text = text.replace(line, '')
        
        # 清理多余空行
        text = re.sub(r'\n{3,}', '\n\n', text)
        
        return text, count
        
    def copy_result(self):
        """复制结果"""
        text = self.output_text.get('1.0', tk.END).strip()
        if text:
            self.root.clipboard_clear()
            self.root.clipboard_append(text)
            self.log("已复制到剪贴板", "success")
            messagebox.showinfo("提示", "已复制到剪贴板！")
        else:
            messagebox.showwarning("警告", "没有可复制的内容！")
            
    def save_result(self):
        """保存结果"""
        text = self.output_text.get('1.0', tk.END).strip()
        if not text:
            messagebox.showwarning("警告", "没有可保存的内容！")
            return
            
        filepath = filedialog.asksaveasfilename(
            title="保存文件",
            defaultextension=".txt",
            filetypes=[
                ("文本文件", "*.txt"),
                ("Word 文档", "*.docx"),
            ]
        )
        
        if filepath:
            try:
                if filepath.endswith('.docx') and HAS_DOCX:
                    doc = docx.Document()
                    for line in text.split('\n'):
                        doc.add_paragraph(line)
                    doc.save(filepath)
                else:
                    with open(filepath, 'w', encoding='utf-8') as f:
                        f.write(text)
                self.log(f"已保存: {filepath}", "success")
            except Exception as e:
                self.log(f"保存失败: {e}", "error")
                
    def browse_tesseract(self):
        """浏览 Tesseract 路径"""
        filepath = filedialog.askopenfilename(title="选择 Tesseract 可执行文件")
        if filepath:
            self.tesseract_path.set(filepath)
            
    def update_stats(self):
        """更新统计信息"""
        total = len(self.files)
        to_delete = sum(1 for f in self.files if f.status == 'to-delete')
        to_keep = sum(1 for f in self.files if f.status == 'to-keep')
        processed = sum(1 for f in self.files if f.status == 'processed')
        
        self.stats_var.set(f"文件: {total} | 待删除: {to_delete} | 待保留: {to_keep} | 已处理: {processed} | 替换: {self.replace_count}")
        
    def log(self, message, level='info'):
        """添加日志"""
        timestamp = datetime.now().strftime('%H:%M:%S')
        self.log_queue.put((timestamp, message, level))
        
    def update_log(self):
        """更新日志显示"""
        try:
            while True:
                timestamp, message, level = self.log_queue.get_nowait()
                log_line = f"[{timestamp}] {message}\n"
                self.log_text.insert(tk.END, log_line, level)
                self.log_text.see(tk.END)
        except queue.Empty:
            pass
        
        self.root.after(100, self.update_log)
        
    def clear_log(self):
        """清空日志"""
        self.log_text.delete('1.0', tk.END)


def main():
    """主函数"""
    root = tk.Tk()
    app = DocumentCleanerApp(root)
    root.mainloop()


if __name__ == '__main__':
    main()
